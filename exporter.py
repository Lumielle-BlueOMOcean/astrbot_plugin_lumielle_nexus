"""XLSX export for completed collection tasks."""

from __future__ import annotations

import json
import re
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import Workbook

if __package__:
    from .core import collection_member_stats, format_local_time
else:
    from core import collection_member_stats, format_local_time

_ILLEGAL_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_FORMULA_PREFIXES = ("=", "+", "-", "@")


def _safe_filename(text: str, fallback: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(text or "")).strip(" .")
    return (cleaned or fallback)[:80]


def _safe_cell_value(value: Any) -> str:
    text = "" if value is None else str(value)
    text = _ILLEGAL_CONTROL_RE.sub("", text)
    if text.lstrip(" \t\r\n").startswith(_FORMULA_PREFIXES):
        return "'" + text
    return text


def _safe_row(values: list[Any]) -> list[str]:
    return [_safe_cell_value(value) for value in values]


def export_collection(
    export_dir: Path,
    task: dict[str, Any],
    entries: list[dict[str, Any]],
    members: list[dict[str, Any]] | None = None,
    *,
    self_id: str | None = None,
    target_ids: list[str] | set[str] | None = None,
    timezone_name: str = "UTC",
    identities: list[dict[str, Any]] | None = None,
) -> Path:
    export_dir = Path(export_dir)
    export_dir.mkdir(parents=True, exist_ok=True)
    payload = json.loads(task.get("payload") or "{}")
    fields = [str(field) for field in payload.get("fields", [])]
    title = payload.get("title") or task.get("group_alias") or task["id"]
    output_path = export_dir / f"{_safe_filename(title, task['id'])}_{task['id']}.xlsx"

    workbook = Workbook()
    result_sheet = workbook.active
    result_sheet.title = "统计结果"
    identity_by_user: dict[str, dict[str, str]] = {}
    identity_verified: dict[tuple[str, str], bool] = {}
    for identity in identities or []:
        user_id = str(identity.get("user_id") or "").strip()
        field_name = str(identity.get("field_name") or "").strip()
        value = str(identity.get("value") or "").strip()
        if not user_id or not field_name or not value:
            continue
        key = (user_id, field_name)
        verified = bool(identity.get("verified"))
        if key not in identity_verified or verified or not identity_verified[key]:
            identity_by_user.setdefault(user_id, {})[field_name] = value
            identity_verified[key] = verified
    has_identity_data = bool(identity_by_user)
    member_by_id = {
        str(member.get("user_id") or "").strip(): member
        for member in (members or [])
        if str(member.get("user_id") or "").strip()
    }
    if has_identity_data:
        output_fields = [field for field in fields if field not in {"姓名", "学号"}]
        result_sheet.append(_safe_row([
            "QQ", "姓名", "学号", "群昵称", *output_fields,
            "提交时间", "最后更新时间", "原始提交",
        ]))
    else:
        output_fields = fields
        result_sheet.append(_safe_row([
            "QQ", "群昵称", *fields, "提交时间", "最后更新时间", "原始提交",
        ]))
    for entry in entries:
        parsed = json.loads(entry.get("parsed_data") or "{}")
        member = member_by_id.get(str(entry.get("sender_id") or ""), {})
        fallback_name = str(
            entry.get("sender_name")
            or member.get("card")
            or member.get("nickname")
            or ""
        )
        identity = identity_by_user.get(str(entry.get("sender_id") or ""), {})
        if has_identity_data:
            values = [
                entry.get("sender_id", ""),
                identity.get("姓名") or parsed.get("姓名", "") or fallback_name,
                identity.get("学号") or parsed.get("学号", ""),
                fallback_name,
                *[parsed.get(field, "") for field in output_fields],
                format_local_time(entry.get("submitted_at"), timezone_name),
                format_local_time(entry.get("updated_at"), timezone_name),
                entry.get("raw_message", ""),
            ]
        else:
            values = [
                entry.get("sender_id", ""),
                fallback_name,
                *[parsed.get(field, "") for field in output_fields],
                format_local_time(entry.get("submitted_at"), timezone_name),
                format_local_time(entry.get("updated_at"), timezone_name),
                entry.get("raw_message", ""),
            ]
        result_sheet.append(_safe_row(values))

    missing_sheet = workbook.create_sheet("未提交成员")
    if members is None:
        missing_sheet.append(_safe_row(["无法获取完整群成员名单，未提交人数不可准确计算。"]))
    else:
        missing_sheet.append(_safe_row(["QQ", "群昵称"]))
        stats = collection_member_stats(
            members, entries, self_id=self_id, target_ids=target_ids,
        )
        for member_id in sorted(stats["missing_ids"]):
            member = stats["eligible_members"][member_id]
            missing_sheet.append(_safe_row([
                member_id,
                member.get("card") or member.get("nickname") or "",
            ]))

    info_sheet = workbook.create_sheet("任务信息")
    info_sheet.append(_safe_row(["项目", "值"]))
    info_rows = [
        ("任务 ID", task.get("id", "")),
        ("标题", title),
        ("群", task.get("group_alias") or task.get("group_id", "")),
        ("群号", task.get("group_id", "")),
        ("创建者", task.get("creator_id", "")),
        ("开始时间", format_local_time(task.get("created_at"), timezone_name)),
        ("结束时间", format_local_time(task.get("finished_at"), timezone_name)),
        ("字段", "、".join(fields)),
        ("总提交人数", len(entries)),
    ]
    if payload.get("target_member_set"):
        info_rows.extend([
            ("目标成员集合", payload["target_member_set"]),
            ("目标快照人数", len(payload.get("target_member_ids") or [])),
        ])
    for row in info_rows:
        info_sheet.append(_safe_row(list(row)))

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            font = copy(cell.font)
            font.bold = True
            cell.font = font
    workbook.save(output_path)
    return output_path
