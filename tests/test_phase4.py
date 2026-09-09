import json
import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace

from core import (
    MODERATION_CONFIRM_TTL_SECONDS,
    TaskManager,
    collection_member_stats,
    member_display_name,
    resolve_member_refs,
    search_group_members,
    validate_moderation_preflight,
)
from exporter import export_collection
from qq_adapter import QQAdapter
from storage import Storage


MEMBERS = [
    {"user_id": "9000", "nickname": "机器人", "card": "机器人", "role": "admin"},
    {"user_id": "1001", "nickname": "张三", "card": "学习委员", "role": "member"},
    {"user_id": "1002", "nickname": "李四", "card": "李四", "role": "admin"},
    {"user_id": "1003", "nickname": "张三", "card": "实验 A 组", "role": "member"},
    {"user_id": "1004", "nickname": "自动助手", "is_robot": True, "role": "member"},
]


class MemberResolutionTests(unittest.TestCase):
    def test_search_is_live_contains_query_and_excludes_bot_and_robot(self):
        results = search_group_members(MEMBERS, "张", limit=20, self_id="9000")
        self.assertEqual([item["user_id"] for item in results], ["1001", "1003"])
        self.assertEqual(member_display_name(results[0]), "学习委员")

    def test_exact_qq_id_wins_and_duplicate_refs_dedupe(self):
        resolved = resolve_member_refs(MEMBERS, ["1001", "学习委员", "1001"], self_id="9000")
        self.assertEqual([item["user_id"] for item in resolved], ["1001"])

    def test_exact_duplicate_nickname_is_ambiguous(self):
        with self.assertRaisesRegex(ValueError, "匹配到多个成员"):
            resolve_member_refs(MEMBERS, ["张三"], self_id="9000")

    def test_bot_and_robot_cannot_resolve(self):
        with self.assertRaisesRegex(ValueError, "未找到成员"):
            resolve_member_refs(MEMBERS, ["9000"], self_id="9000")
        with self.assertRaisesRegex(ValueError, "未找到成员"):
            resolve_member_refs(MEMBERS, ["1004"], self_id="9000")

    def test_moderation_role_preflight(self):
        self.assertIsNone(validate_moderation_preflight(
            {"user_id": "9000", "role": "owner"},
            {"user_id": "1002", "role": "admin"},
            "9000", "1002",
        ))
        self.assertIn("管理员", validate_moderation_preflight(
            {"user_id": "9000", "role": "admin"},
            {"user_id": "1002", "role": "admin"},
            "9000", "1002",
        ) or "")
        self.assertIn("群主", validate_moderation_preflight(
            {"user_id": "9000", "role": "owner"},
            {"user_id": "1002", "role": "owner"},
            "9000", "1002",
        ) or "")

    def test_moderation_ttl_is_ten_minutes(self):
        self.assertEqual(MODERATION_CONFIRM_TTL_SECONDS, 600)

    def test_resolution_does_not_use_partial_match(self):
        with self.assertRaisesRegex(ValueError, "未找到成员"):
            resolve_member_refs(MEMBERS, ["学习"], self_id="9000")

    def test_exact_qq_id_beats_a_card_match(self):
        members = [
            {"user_id": "1001", "nickname": "甲", "card": "乙"},
            {"user_id": "乙", "nickname": "丙", "card": "丁"},
        ]
        self.assertEqual(resolve_member_refs(members, ["乙"])[0]["user_id"], "乙")

    def test_moderation_preflight_rejects_non_admin_bot(self):
        error = validate_moderation_preflight(
            {"user_id": "9000", "role": "member"},
            {"user_id": "1001", "role": "member"},
            "9000", "1001",
        )
        self.assertIn("不是该群管理员", error or "")


class MemberSetAndCollectionTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.storage = Storage(Path(self.temp_dir.name))
        self.manager = TaskManager(self.storage)
        await self.manager.bind_group("班群", "123456789", "qq-main", "operator")

    async def asyncTearDown(self):
        self.storage.close()
        self.temp_dir.cleanup()

    async def test_member_set_tables_are_present_with_foreign_keys_enabled(self):
        tables = {
            row[0] for row in self.storage._conn.execute(
                "SELECT name FROM sqlite_master WHERE type = 'table'",
            ).fetchall()
        }
        self.assertIn("member_sets", tables)
        self.assertIn("member_set_members", tables)
        self.assertEqual(self.storage._conn.execute("PRAGMA foreign_keys").fetchone()[0], 1)

    async def test_schema_member_sets_are_group_scoped_and_update_modes_work(self):
        await self.manager.set_member_set(
            "班群", "班委", ["1001", "1002", "1002"], "replace",
            "qq-main", "operator", MEMBERS, "9000",
        )
        await self.manager.bind_group("社团群", "987654321", "qq-main", "operator")
        await self.manager.set_member_set(
            "社团群", "班委", ["1003"], "replace",
            "qq-main", "operator", MEMBERS, "9000",
        )
        await self.manager.set_member_set(
            "班群", "班委", ["1003"], "add",
            "qq-main", "operator", MEMBERS, "9000",
        )
        await self.manager.set_member_set(
            "班群", "班委", ["1002"], "remove",
            "qq-main", "operator", MEMBERS, "9000",
        )
        first = await self.manager.get_member_set("班群", "班委", "qq-main", MEMBERS, "9000")
        second = await self.manager.get_member_set("社团群", "班委", "qq-main", MEMBERS, "9000")
        self.assertEqual({item["user_id"] for item in first["members"]}, {"1001", "1003"})
        self.assertEqual({item["user_id"] for item in second["members"]}, {"1003"})

    async def test_member_set_rejects_empty_and_preserves_departed_member(self):
        await self.manager.set_member_set(
            "班群", "实验A组", ["1001", "1002"], "replace",
            "qq-main", "operator", MEMBERS, "9000",
        )
        with self.assertRaisesRegex(ValueError, "不能为空"):
            await self.manager.set_member_set(
                "班群", "实验A组", [], "replace",
                "qq-main", "operator", MEMBERS, "9000",
            )
        live_without_1002 = [member for member in MEMBERS if member["user_id"] != "1002"]
        found = await self.manager.get_member_set(
            "班群", "实验A组", "qq-main", live_without_1002, "9000",
        )
        self.assertFalse(found["members_by_id"]["1002"]["present"])

    async def test_member_set_caps_at_two_thousand_members(self):
        many_members = [
            {"user_id": str(index), "nickname": f"成员{index}"}
            for index in range(2100)
        ]
        with self.assertRaisesRegex(ValueError, "2000"):
            await self.manager.set_member_set(
                "班群", "超大名单", [str(index) for index in range(2100)], "replace",
                "qq-main", "operator", many_members, "9000",
            )

    async def test_target_collection_snapshots_set_and_ignores_non_target(self):
        await self.manager.set_member_set(
            "班群", "班委", ["1001", "1002"], "replace",
            "qq-main", "operator", MEMBERS, "9000",
        )
        task = await self.manager.start_collection(
            "班群", "报名", ["姓名"], "", False,
            "qq-main", "operator", "origin", "班委",
        )
        payload = json.loads(task["payload"])
        self.assertEqual(payload["target_member_ids"], ["1001", "1002"])
        self.assertIsNone(await self.manager.process_collection_message(
            "qq-main", "123456789", "1003", "张三", "姓名：张三",
        ))
        accepted = await self.manager.process_collection_message(
            "qq-main", "123456789", "1001", "张三", "姓名：张三",
        )
        self.assertIsNotNone(accepted)
        await self.manager.set_member_set(
            "班群", "班委", ["1003"], "replace",
            "qq-main", "operator", MEMBERS, "9000",
        )
        unchanged = json.loads(self.storage.get_task(task["id"])["payload"])
        self.assertEqual(unchanged["target_member_ids"], ["1001", "1002"])

    async def test_target_stats_intersects_current_human_members(self):
        entries = [{"sender_id": "1001"}, {"sender_id": "1003"}]
        live = [member for member in MEMBERS if member["user_id"] != "1002"]
        stats = collection_member_stats(live, entries, "9000", ["1001", "1002"])
        self.assertEqual(stats["eligible_ids"], {"1001"})
        self.assertEqual(stats["missing_ids"], set())


class RelayAndModerationStorageTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.storage = Storage(Path(self.temp_dir.name))
        self.manager = TaskManager(self.storage)
        await self.manager.bind_group("班群", "123456789", "qq-main", "operator")

    async def asyncTearDown(self):
        self.storage.close()
        self.temp_dir.cleanup()

    async def test_relay_member_snapshot_is_persisted_and_mutual_exclusion_rejected(self):
        await self.manager.set_member_set(
            "班群", "班委", ["1001", "1002"], "replace",
            "qq-main", "operator", MEMBERS, "9000",
        )
        relay = await self.manager.prepare_relay(
            "班群", "请查看", False, "", "qq-main", "operator", "origin",
            "班委", ["1001", "1002"],
        )
        payload = json.loads(relay["payload"])
        self.assertEqual(payload["mention_member_set"], "班委")
        self.assertEqual(payload["mention_user_ids"], ["1001", "1002"])
        with self.assertRaisesRegex(ValueError, "互斥"):
            await self.manager.prepare_relay(
                "班群", "请查看", True, "", "qq-main", "operator", "origin",
                "班委", ["1001"],
            )

    async def test_relay_member_snapshot_survives_set_changes(self):
        await self.manager.set_member_set(
            "班群", "班委", ["1001", "1002"], "replace",
            "qq-main", "operator", MEMBERS, "9000",
        )
        relay = await self.manager.prepare_relay(
            "班群", "请查看", False, "", "qq-main", "operator", "origin",
            "班委", ["1001", "1002"],
        )
        await self.manager.set_member_set(
            "班群", "班委", ["1003"], "replace",
            "qq-main", "operator", MEMBERS, "9000",
        )
        self.assertEqual(json.loads(self.storage.get_task(relay["id"])["payload"])["mention_user_ids"], ["1001", "1002"])

    async def test_moderation_claim_requires_creator_and_ttl_and_never_retries(self):
        task = await self.manager.prepare_moderation(
            "班群", "mute", "1001", "张三", "member", "admin", 600,
            False, "刷屏", "qq-main", "operator", "origin",
        )
        with self.assertRaisesRegex(ValueError, "只能由创建"):
            await self.manager.confirm_moderation(task["id"], "qq-main", "other")
        claimed = await self.manager.confirm_moderation(task["id"], "qq-main", "operator")
        self.assertEqual(claimed["status"], "PROCESSING")
        failed = await self.manager.finish_moderation(task["id"], False, "adapter failed")
        self.assertEqual(failed["status"], "FAILED")
        self.assertEqual(failed["retry_count"], 0)

    async def test_moderation_processing_recovery_is_unknown_outcome(self):
        task = await self.manager.prepare_moderation(
            "班群", "kick", "1001", "张三", "member", "owner", 0,
            True, "", "qq-main", "operator", "origin",
        )
        await self.manager.confirm_moderation(task["id"], "qq-main", "operator")
        self.storage.close()
        self.storage = Storage(Path(self.temp_dir.name))
        TaskManager(self.storage)
        recovered = self.storage.get_task(task["id"])
        self.assertEqual(recovered["status"], "FAILED")
        self.assertEqual(json.loads(recovered["result"])["execution_outcome"], "unknown")

    async def test_expired_moderation_preview_is_cancelled(self):
        task = await self.manager.prepare_moderation(
            "班群", "mute", "1001", "张三", "member", "owner", 600,
            False, "", "qq-main", "operator", "origin",
        )
        with self.assertRaisesRegex(ValueError, "preview 已过期"):
            await self.manager.confirm_moderation(
                task["id"], "qq-main", "operator",
                now=datetime.now(timezone.utc) + timedelta(seconds=601),
            )
        self.assertEqual(self.storage.get_task(task["id"])["status"], "CANCELLED")


class AdapterPhase4Tests(unittest.IsolatedAsyncioTestCase):
    async def test_member_info_and_mutation_actions_use_onebot_api(self):
        calls = []

        class Client:
            async def call_action(self, action, **kwargs):
                calls.append((action, kwargs))
                if action == "get_group_member_info":
                    return {"user_id": kwargs["user_id"], "role": "member"}
                return {}

        class Platform:
            def meta(self):
                return SimpleNamespace(id="qq-main", name="aiocqhttp")

            def get_client(self):
                return Client()

        class Context:
            def get_platform_inst(self, platform_id):
                return Platform()

        adapter = QQAdapter(Context(), "qq-main")
        info = await adapter.get_group_member_info("123456789", "1001")
        await adapter.set_group_ban("123456789", "1001", 600)
        await adapter.set_group_kick("123456789", "1001", True)
        self.assertEqual(info["user_id"], 1001)
        self.assertEqual([call[0] for call in calls], [
            "get_group_member_info", "set_group_ban", "set_group_kick",
        ])
        self.assertEqual(calls[1][1]["duration"], 600)
        self.assertTrue(calls[2][1]["reject_add_request"])


class TargetCollectionExportTests(unittest.TestCase):
    def test_target_snapshot_limits_missing_sheet_and_info(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            task = {
                "id": "C-20260909-004",
                "group_alias": "班群",
                "group_id": "123456789",
                "creator_id": "operator",
                "created_at": "2026-09-09T01:00:00+00:00",
                "finished_at": "2026-09-09T02:00:00+00:00",
                "payload": json.dumps({
                    "title": "定向统计",
                    "fields": ["内容"],
                    "target_member_set": "班委",
                    "target_member_ids": ["20001", "20002"],
                }, ensure_ascii=False),
            }
            output = export_collection(
                Path(temp_dir), task, [], [
                    {"user_id": "20001", "nickname": "张三"},
                    {"user_id": "20002", "nickname": "李四"},
                    {"user_id": "20003", "nickname": "王五"},
                ], target_ids=["20001", "20002"], self_id="99999",
            )
            from openpyxl import load_workbook

            workbook = load_workbook(output, read_only=True)
            missing = list(workbook["未提交成员"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(missing, [("20001", "张三"), ("20002", "李四")])
            info = list(workbook["任务信息"].iter_rows(values_only=True))
            self.assertIn(("目标成员集合", "班委"), info)
            self.assertIn(("目标快照人数", "2"), info)
