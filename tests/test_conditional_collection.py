import json
import importlib
import sys
import tempfile
import types
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest import mock

import core as core_module
from core import Storage, TaskManager, collection_member_stats, parse_collection_submission
from exporter import export_collection


FIELDS = ["返校情况", "返校时间", "原因"]
MAPPINGS = {
    "返校情况": {
        "已返校": ["1", "已返校", "到了"],
        "未返校": ["2", "未返校", "还没回"],
    },
}
REQUIRED_WHEN = {
    "返校时间": {"field": "返校情况", "equals": ["未返校"]},
    "原因": {"field": "返校情况", "equals": ["未返校"]},
}


def _load_main_module():
    try:
        return importlib.import_module("main")
    except ModuleNotFoundError as exc:
        if exc.name != "astrbot":
            raise
        try:
            from test_collection_incident import _load_main_module as load_with_astrbot_stubs
        except ModuleNotFoundError:
            from tests.test_collection_incident import _load_main_module as load_with_astrbot_stubs

        return load_with_astrbot_stubs()


class ConditionalCollectionTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.storage = Storage(Path(self.temp_dir.name))
        self.manager = TaskManager(self.storage, timezone_name="Asia/Shanghai")
        await self.manager.bind_group("班群", "123456789", "qq-main", "operator")
        self.now = datetime.now(timezone.utc).replace(microsecond=0)

    async def asyncTearDown(self):
        self.storage.close()
        self.temp_dir.cleanup()

    async def _collection(self, **kwargs):
        return await self.manager.start_collection(
            "班群", "返校统计", FIELDS, "", False,
            "qq-main", "operator", "private:operator",
            now=self.now, field_value_mappings=MAPPINGS, **kwargs,
        )

    def test_conditional_completion_helper_uses_canonical_values(self):
        required_fields = getattr(core_module, "collection_required_fields", None)
        completion = getattr(core_module, "collection_entry_completion", None)
        self.assertTrue(callable(required_fields), "single required-field helper is missing")
        self.assertTrue(callable(completion), "single completion helper is missing")
        self.assertEqual(
            required_fields(FIELDS, {"返校情况": "已返校"}, REQUIRED_WHEN),
            {"返校情况"},
        )
        self.assertEqual(
            required_fields(FIELDS, {"返校情况": "未返校"}, REQUIRED_WHEN),
            set(FIELDS),
        )
        self.assertEqual(
            completion(FIELDS, {"返校情况": "已返校"}, REQUIRED_WHEN),
            "COMPLETE",
        )
        self.assertEqual(
            completion(FIELDS, {"返校情况": "未返校", "返校时间": "10月9日"}, REQUIRED_WHEN),
            "PARTIAL",
        )
        self.assertEqual(
            completion(FIELDS, {
                "返校情况": "未返校", "返校时间": "10月9日", "原因": "家中有事",
            }, REQUIRED_WHEN),
            "COMPLETE",
        )

    def test_stats_share_conditional_complete_partial_and_missing_semantics(self):
        stats = collection_member_stats(
            [{"user_id": str(user_id)} for user_id in (1001, 1002, 1003)],
            [
                {"sender_id": "1001", "parsed_data": '{"返校情况":"已返校"}'},
                {"sender_id": "1002", "parsed_data": '{"返校情况":"未返校"}'},
            ],
            required_fields=FIELDS,
            required_when=REQUIRED_WHEN,
        )
        self.assertEqual(stats["complete_ids"], {"1001"})
        self.assertEqual(stats["partial_ids"], {"1002"})
        self.assertEqual(stats["no_response_ids"], {"1003"})
        self.assertEqual(stats["missing_ids"], {"1002", "1003"})

    def test_missing_required_when_preserves_legacy_all_fields_required(self):
        stats = collection_member_stats(
            [{"user_id": "1001"}],
            [{"sender_id": "1001", "parsed_data": '{"返校情况":"已返校"}'}],
            required_fields=FIELDS,
        )
        self.assertEqual(stats["complete_ids"], set())
        self.assertEqual(stats["partial_ids"], {"1001"})

    def test_multifield_bare_unique_mapping_is_deterministic(self):
        self.assertEqual(
            parse_collection_submission(FIELDS, "1", MAPPINGS),
            {"返校情况": "已返校"},
        )

    def test_multifield_bare_ambiguous_mapping_fails_closed(self):
        self.assertEqual(
            parse_collection_submission(
                ["状态", "等级"], "1",
                {"状态": {"已返校": ["1"]}, "等级": {"优秀": ["1"]}},
            ),
            {},
        )
        self.assertEqual(
            parse_collection_submission(
                ["状态", "等级"], "1",
                {
                    "状态": {"甲": ["1"], "乙": ["1"]},
                    "等级": {"优秀": ["1"]},
                },
            ),
            {},
        )

    def test_structured_field_resolves_reused_alias_without_bare_ambiguity(self):
        self.assertEqual(
            parse_collection_submission(
                ["返校情况", "原因"], "返校情况：1",
                {"返校情况": {"已返校": ["1"]}, "原因": {"一级": ["1"]}},
            ),
            {"返校情况": "已返校"},
        )

    async def test_required_when_is_validated_and_persisted_in_task_payload(self):
        task = await self._collection(required_when=REQUIRED_WHEN)
        self.assertEqual(json.loads(task["payload"])["required_when"], REQUIRED_WHEN)

        invalid_rules = [
            {"不存在": {"field": "返校情况", "equals": ["未返校"]}},
            {"原因": {"field": "不存在", "equals": ["未返校"]}},
            {"返校情况": {"field": "返校情况", "equals": ["未返校"]}},
            {"原因": {"field": "返校情况", "equals": []}},
            {"原因": {"field": "返校情况", "equals": "未返校"}},
            {"原因": {"field": "返校情况", "equals": [""]}},
            {
                "返校时间": {"field": "返校情况", "equals": ["未返校"]},
                "原因": {"field": "返校情况", "equals": ["未返校"]},
                "返校情况": {"field": "原因", "equals": ["家中有事"]},
            },
        ]
        for rules in invalid_rules:
            with self.subTest(rules=rules), self.assertRaises(ValueError):
                await self._collection(required_when=rules)

    async def test_bare_reply_is_persisted_deterministically_and_conditionally_complete(self):
        task = await self._collection(
            required_when=REQUIRED_WHEN, ai_extraction=True,
            ai_provider_id="provider-1",
        )
        result = await self.manager.capture_collection_message(
            task["id"], "1001", "张三", "1", now=self.now,
            source_message_id="bare-returned",
        )
        self.assertTrue(result["deterministic_handled"])
        self.assertEqual(result["entry"]["parsed_data"], {"返校情况": "已返校"})
        self.assertEqual(result["missing"], [])
        stats = collection_member_stats(
            [{"user_id": "1001"}], [self.storage.get_entry(task["id"], "1001")],
            required_fields=FIELDS, required_when=REQUIRED_WHEN,
        )
        self.assertEqual(stats["complete_ids"], {"1001"})
        ai_result = await self.manager.apply_collection_ai_candidate(
            {
                "task_id": task["id"], "platform_id": "qq-main",
                "group_id": "123456789", "sender_id": "1002",
                "body": "我已经返校", "mode": "fill",
            },
            [{
                "field": "返校情况", "value": "已返校",
                "evidence": "已经返校", "confidence": 0.99,
            }],
            "我已经返校", "李四",
        )
        self.assertEqual(ai_result["missing"], [])

    async def test_status_fallback_counts_only_complete_entries(self):
        task = await self._collection(required_when=REQUIRED_WHEN)
        await self.manager.capture_collection_message(
            task["id"], "1001", "张三", "1", now=self.now,
            source_message_id="complete-returned",
        )
        await self.manager.capture_collection_message(
            task["id"], "1002", "李四", "2", now=self.now,
            source_message_id="partial-not-returned",
        )
        status = await self.manager.collection_status(task["id"], "qq-main")
        self.assertEqual(status["submitted_count"], 2)
        self.assertEqual(status["complete_count"], 1)

    async def test_conditionally_complete_member_is_not_no_data_recheck_candidate(self):
        task = await self._collection(required_when=REQUIRED_WHEN, ai_extraction=True,
                                      ai_provider_id="provider-1")
        await self.manager.capture_collection_message(
            task["id"], "1001", "张三", "1", now=self.now,
            source_message_id="bare-returned",
        )
        evidence = await self.manager.capture_collection_message(
            task["id"], "1001", "张三", "我还要补充一句", now=self.now,
            source_message_id="later-evidence",
        )
        self.storage.disposition_workflow_messages(
            [evidence["workflow_message"]["id"]], {"1001": "no_data"}, {},
            self.now.isoformat(),
        )
        snapshot = await self.manager.prepare_collection_checkpoint(
            task["id"], self.now, "chase",
        )
        self.assertEqual(snapshot["messages"], [])

    async def test_real_chase_mentions_only_partial_and_no_response_members(self):
        main_module = _load_main_module()
        collection = await self._collection(required_when=REQUIRED_WHEN)
        first = await self.manager.capture_collection_message(
            collection["id"], "1001", "张三", "1", now=self.now,
            source_message_id="returned",
        )
        second = await self.manager.capture_collection_message(
            collection["id"], "1002", "李四", "2", now=self.now,
            source_message_id="not-returned",
        )
        self.assertTrue(first["deterministic_handled"])
        self.assertTrue(second["deterministic_handled"])
        chase = await self.manager.schedule_collection_chase(
            collection["id"], self.now + timedelta(minutes=1), "请补充", 0,
            "qq-main", "operator", "private:operator",
        )
        chase_payload = json.loads(chase["payload"])

        class FakeAdapter:
            def __init__(self):
                self.sent = []

            async def get_group_member_list(self, _group_id):
                return [
                    {"user_id": "1001", "nickname": "张三"},
                    {"user_id": "1002", "nickname": "李四"},
                    {"user_id": "1003", "nickname": "王五"},
                    {"user_id": "9000", "nickname": "Bot", "is_robot": True},
                ]

            async def send_group_at_member_batch(self, group_id, user_ids, text):
                self.sent.append((group_id, list(user_ids), text))

        adapter = FakeAdapter()
        plugin = object.__new__(main_module.LumielleNexus)
        plugin.context = object()
        plugin.manager = self.manager
        plugin.storage = self.storage
        plugin._reconcile_collection_history = mock.AsyncMock(return_value=(
            adapter, "9000", {"coverage_status": "FULL"},
        ))
        plugin._run_collection_checkpoint = mock.AsyncMock(return_value={
            "analysis_incomplete": False,
        })

        await plugin._execute_collection_chase(chase, chase_payload)
        self.assertEqual(len(adapter.sent), 1)
        self.assertEqual(adapter.sent[0][1], ["1002", "1003"])

        await self.manager.capture_collection_message(
            collection["id"], "1002", "李四",
            "返校时间：10月9日上午\n原因：家中有事", now=self.now,
            source_message_id="not-returned-details",
        )
        second_chase = await self.manager.schedule_collection_chase(
            collection["id"], self.now + timedelta(minutes=2), "请补充", 0,
            "qq-main", "operator", "private:operator",
        )
        await plugin._execute_collection_chase(
            second_chase, json.loads(second_chase["payload"]),
        )
        self.assertEqual(len(adapter.sent), 2)
        self.assertEqual(adapter.sent[1][1], ["1003"])

    async def test_manual_finalize_defaults_and_exports_use_conditional_completeness(self):
        main_module = _load_main_module()
        collection = await self._collection(
            required_when=REQUIRED_WHEN,
            missing_default_field="返校情况",
            missing_default_value="未返校",
            deadline=self.now + timedelta(hours=2),
            auto_export=True,
        )
        await self.manager.capture_collection_message(
            collection["id"], "1001", "张三", "1", now=self.now,
            source_message_id="returned",
        )
        await self.manager.capture_collection_message(
            collection["id"], "1002", "李四", "2", now=self.now,
            source_message_id="not-returned",
        )
        await self.manager.capture_collection_message(
            collection["id"], "1002", "李四",
            "返校时间：10月9日上午\n原因：家中有事", now=self.now,
            source_message_id="not-returned-details",
        )
        stopped = await self.manager.stop_collection(collection["id"], "qq-main")
        cutoff = json.loads(stopped["task"]["result"])["processing_cutoff"]

        class FakeAdapter:
            async def get_group_member_list(self, _group_id):
                return [
                    {"user_id": "1001", "nickname": "张三"},
                    {"user_id": "1002", "nickname": "李四"},
                    {"user_id": "1003", "nickname": "王五"},
                    {"user_id": "9000", "nickname": "Bot", "is_robot": True},
                ]

            async def upload_private_file(self, *_args):
                return None

        adapter = FakeAdapter()
        plugin = object.__new__(main_module.LumielleNexus)
        plugin.manager = self.manager
        plugin.storage = self.storage
        plugin.data_dir = Path(self.temp_dir.name)
        plugin._reconcile_collection_history = mock.AsyncMock(return_value=(
            adapter, "9000", {"coverage_status": "FULL"},
        ))
        plugin._run_collection_checkpoint = mock.AsyncMock(return_value={
            "analysis_incomplete": False,
        })

        await plugin._execute_collection_finalize(
            {"id": "manual-finalize", "platform_id": "qq-main", "run_at": cutoff},
            {"collection_task_id": collection["id"], "manual_stop": True,
             "force_export": True, "scheduled_run_at": cutoff},
        )
        final_task = self.storage.get_task(collection["id"])
        self.assertEqual(final_task["status"], "COMPLETED")
        entries = {row["sender_id"]: json.loads(row["parsed_data"])
                   for row in self.storage.list_entries(collection["id"])}
        self.assertEqual(entries["1001"], {"返校情况": "已返校"})
        self.assertEqual(entries["1002"], {
            "返校情况": "未返校", "返校时间": "10月9日上午", "原因": "家中有事",
        })
        self.assertEqual(entries["1003"], {"返校情况": "未返校"})

        from openpyxl import load_workbook

        workbook = load_workbook(json.loads(final_task["result"])["export_path"],
                                 read_only=True, data_only=True)
        try:
            missing_rows = list(workbook["未提交成员"].iter_rows(values_only=True))
            missing_ids = {str(row[0]) for row in missing_rows[1:] if row[0]}
            self.assertNotIn("1001", missing_ids)
            self.assertNotIn("1002", missing_ids)
            self.assertIn("1003", missing_ids)
        finally:
            workbook.close()

    def test_collection_ack_defaults_off_and_respects_explicit_values(self):
        main_module = _load_main_module()
        schema = json.loads(Path(__file__).resolve().parents[1].joinpath(
            "_conf_schema.json",
        ).read_text())
        self.assertFalse(schema["collection_ack"]["default"])
        for config, expected in (({}, False), ({"collection_ack": False}, False),
                                 ({"collection_ack": True}, True)):
            with self.subTest(config=config), tempfile.TemporaryDirectory() as directory:
                with mock.patch.object(
                    main_module.StarTools, "get_data_dir", return_value=Path(directory),
                ):
                    plugin = main_module.LumielleNexus(object(), config)
                try:
                    self.assertIs(plugin.collection_ack, expected)
                finally:
                    plugin.storage.close()

    async def test_ack_setting_changes_reply_but_never_deterministic_capture(self):
        main_module = _load_main_module()
        collection = await self._collection(required_when=REQUIRED_WHEN)
        event = types.SimpleNamespace(
            message_obj=types.SimpleNamespace(timestamp=self.now.timestamp(), message_id="ack-1"),
            get_sender_id=lambda: "1001",
            get_sender_name=lambda: "张三",
            get_self_id=lambda: "9000",
            get_group_id=lambda: "123456789",
            get_platform_id=lambda: "qq-main",
            get_message_str=lambda: "1",
            plain_result=lambda value: value,
        )
        plugin = object.__new__(main_module.LumielleNexus)
        plugin.manager = self.manager
        plugin.storage = self.storage
        plugin._handle_poll_message = mock.AsyncMock(return_value={"handled": False})
        plugin.collection_ack = False
        silent = [item async for item in plugin.on_group_message(event)]
        self.assertEqual(silent, [])
        self.assertEqual(
            json.loads(self.storage.get_entry(collection["id"], "1001")["parsed_data"]),
            {"返校情况": "已返校"},
        )

        event.message_obj.message_id = "ack-2"
        plugin.collection_ack = True
        acknowledged = [item async for item in plugin.on_group_message(event)]
        self.assertEqual(acknowledged, ["已记录。"])
        self.assertIsNotNone(self.storage.get_entry(collection["id"], "1001"))

    async def test_collection_announcement_explains_mapping_and_conditional_fields(self):
        main_module = _load_main_module()

        class Event:
            unified_msg_origin = "aiocqhttp:FriendMessage:operator"

            def is_private_chat(self):
                return True

            def is_admin(self):
                return True

            def get_platform_name(self):
                return "aiocqhttp"

            def get_platform_id(self):
                return "qq-main"

            def get_sender_id(self):
                return "operator"

        class FakeAdapter:
            def __init__(self):
                self.messages = []

            async def send_group_text(self, group_id, text):
                self.messages.append((group_id, text))

        adapter = FakeAdapter()
        plugin = main_module.LumielleNexus.__new__(main_module.LumielleNexus)
        plugin.manager = self.manager
        plugin.storage = self.storage
        plugin.operator_ids = set()
        plugin._adapter = lambda _event: adapter

        result = await plugin._start_collection(
            Event(), "班群", "返校统计", FIELDS, "", False,
            field_value_mappings=MAPPINGS, required_when=REQUIRED_WHEN,
        )
        self.assertIn("已开始收集", result)
        notice = adapter.messages[0][1]
        self.assertIn("1 = 已返校", notice)
        self.assertIn("2 = 未返校", notice)
        self.assertIn("返校时间：", notice)
        self.assertIn("原因：", notice)
        self.assertLessEqual(len(notice), 1200)

    def test_export_missing_sheet_omits_conditionally_complete_member(self):
        with tempfile.TemporaryDirectory() as export_dir:
            task = {
                "id": "C-20261007-001", "group_alias": "班群", "group_id": "123456789",
                "creator_id": "operator", "created_at": self.now.isoformat(),
                "finished_at": self.now.isoformat(),
                "payload": json.dumps({
                    "title": "返校统计", "fields": FIELDS, "required_when": REQUIRED_WHEN,
                }, ensure_ascii=False),
            }
            path = export_collection(
                Path(export_dir), task,
                [
                    {"sender_id": "1001", "sender_name": "张三",
                     "parsed_data": '{"返校情况":"已返校"}'},
                    {"sender_id": "1002", "sender_name": "李四",
                     "parsed_data": '{"返校情况":"未返校"}'},
                ],
                [{"user_id": "1001", "nickname": "张三"},
                 {"user_id": "1002", "nickname": "李四"}],
                timezone_name="Asia/Shanghai",
            )
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                missing_rows = list(workbook["未提交成员"].iter_rows(values_only=True))
                missing_ids = {str(row[0]) for row in missing_rows[1:] if row[0]}
                self.assertNotIn("1001", missing_ids)
                self.assertIn("1002", missing_ids)
                result_rows = list(workbook["统计结果"].iter_rows(values_only=True))
                self.assertIn("已返校", {str(cell) for row in result_rows for cell in row})
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
